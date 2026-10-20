"""request_template.py — the data request a consultant sends a client, generated from the canonical schema.

    request_rows(params=None) -> [dict]   one per field: table, field, required, what it is, unit / format
    request_html(params=None) -> str      one printable page
    request_xlsx(params=None) -> bytes    the same as a spreadsheet

Built from schema.TABLES (required flags), FIELD_LABEL / FIELD_HELP, the
period wanted (IngestParams coverage settings) and the file formats ingest
reads (readers.SUPPORTED), so it can never ask for something ingest ignores.
"""
from __future__ import annotations

import calendar
import io
from html import escape

from .params import IngestParams
from .readers import SUPPORTED
from .schema import FIELD_HELP, FIELD_LABEL, REQUIRED_TABLES, TABLES, TABLE_LABEL

_FORMAT = {"id": "text", "text": "text", "length": "inches or cm (say which)", "weight": "lb or kg (say which)",
           "count": "units per case", "qty": "units (eaches)", "cases": "number of full cases",
           "stock": "units on hand", "date": "a date (any common format)", "time": "time of day"}
_WHY = {"item_master": "container fit, cube and storage",
        "orders": "volumes, the design day, picks per hour and workstations",
        "receipts": "inbound and decant sizing", "inventory": "storage capacity and days on hand"}


def _period(params: IngestParams) -> str:
    peak = " and ".join(calendar.month_name[m] for m in params.peak_months)
    return (f"At least 12 months of history ({params.min_coverage_days}+ days) ending as recently as possible, "
            f"including the peak season ({peak}).")


def request_rows(params: IngestParams | None = None) -> list[dict]:
    rows = []
    for table, flds in TABLES.items():
        for f in flds:
            rows.append(dict(table=TABLE_LABEL[table], table_needed="Required" if table in REQUIRED_TABLES else
                             "Optional", field=FIELD_LABEL[f.name], required="Required" if f.required else "Optional",
                             what=FIELD_HELP[table].get(f.name, ""), unit=_FORMAT.get(f.kind, ""),
                             also_called=", ".join(s for s in f.synonyms[:3])))
    return rows


def request_html(params: IngestParams | None = None) -> str:
    params = params or IngestParams()
    rows = request_rows(params)
    css = ("body{font:14px/1.45 system-ui,Segoe UI,sans-serif;color:#1d2433;max-width:900px;margin:24px auto;"
           "padding:0 16px}h1{font-size:22px;margin:0}h2{font-size:16px;margin:20px 0 6px}table{border-collapse:"
           "collapse;width:100%;font-size:13px}th,td{border-bottom:1px solid #d9dee7;padding:4px 6px;text-align:left;"
           "vertical-align:top}.req{font-weight:600}.muted{color:#5b6475}@media print{body{margin:0}}")
    h = [f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>Data request</title>"
         f"<style>{css}</style></head><body>",
         "<h1>Warehouse data request</h1>",
         "<p>To size a goods-to-person system we need the four extracts below, exported straight from your WMS or "
         "ERP. One row per item, order line, receipt line or stock record; no totals, pivots or formatting.</p>",
         f"<p><b>Period:</b> {escape(_period(params))}</p>",
         "<p><b>Units:</b> sizes in inches or centimetres and weights in pounds or kilograms; please say which "
         "(a unit in the column name is enough, e.g. <i>Length (cm)</i>). Quantities in units (eaches); if your "
         "export counts full cases separately, include both columns.</p>"]
    for label in dict.fromkeys(r["table"] for r in rows):
        t_rows = [r for r in rows if r["table"] == label]
        table_key = next(k for k, v in TABLE_LABEL.items() if v == label)
        h.append(f"<h2>{escape(label)} <span class='muted'>({t_rows[0]['table_needed'].lower()}: "
                 f"used for {_WHY[table_key]})</span></h2>")
        h.append("<table><tr><th>Field</th><th></th><th>What it is</th><th>Unit / format</th><th>Often called</th></tr>")
        for r in t_rows:
            cls = " class='req'" if r["required"] == "Required" else ""
            h.append(f"<tr><td{cls}>{escape(r['field'])}</td><td>{r['required']}</td><td>{escape(r['what'])}</td>"
                     f"<td>{escape(r['unit'])}</td><td class='muted'>{escape(r['also_called'])}</td></tr>")
        h.append("</table>")
    h.append("<h2>File formats we can read</h2><ul>")
    h += [f"<li><b>{escape(k)}</b>: {escape(v)}</li>" for k, v in SUPPORTED]
    h.append("</ul><p class='muted'>Excel files are fine, but a sheet stops at 1,048,576 rows; for larger "
             "extracts please send CSV or text.</p></body></html>")
    return "\n".join(h)


def request_xlsx(params: IngestParams | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    params = params or IngestParams()
    wb = Workbook()
    ws = wb.active
    ws.title = "Data request"
    ws.append(["Warehouse data request"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([_period(params)])
    ws.append(["Units: sizes in inches or cm, weights in lb or kg (say which); quantities in units (eaches)."])
    ws.append([])
    head = ["Table", "Table needed", "Field", "Field needed", "What it is", "Unit / format", "Often called"]
    ws.append(head)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for r in request_rows(params):
        ws.append([r["table"], r["table_needed"], r["field"], r["required"], r["what"], r["unit"], r["also_called"]])
    for col, width in zip("ABCDEFG", (22, 12, 30, 12, 70, 26, 40)):
        ws.column_dimensions[col].width = width
    fmt = wb.create_sheet("File formats")
    fmt.append(["Format", "Details"])
    for k, v in SUPPORTED:
        fmt.append([k, v])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
