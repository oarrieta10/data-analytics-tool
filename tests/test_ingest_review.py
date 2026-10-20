"""Backend for the guided Ingest page: detection, coverage, plain-language issues, reasoned acceptance,
progress, and the client-facing outputs (receipt, rejects, data request template).

All on the frozen sample client in tests/fixtures (no real client data).
"""
import gzip
import re
import sqlite3
import zipfile
from pathlib import Path

import duckdb
import pandas as pd
import pytest

from gtpfit.ingest import IngestParams, ingest_folder, load_wms_profile, write_outputs
from gtpfit.ingest.checks import ISSUE_TEXT
from gtpfit.ingest.coverage import coverage
from gtpfit.ingest.detect import detect

ROOT = Path(__file__).resolve().parents[1]
FIX = ROOT / "tests" / "fixtures"
SAMPLE = FIX / "sample_client"
WMS = FIX / "sample_wms_profile.yaml"
FILES = {"item_master": "item_master", "orders": "orders", "receipts": "receipts", "inventory": "inventory_snapshots"}


def _raw(name):
    return pd.read_csv(SAMPLE / f"{name}.csv", dtype=str, keep_default_na=False)


def _typed(name):
    return duckdb.sql(f"select * from read_csv('{(SAMPLE / (name + '.csv')).as_posix()}', header=true)").df()


def _clone(dst: Path, **frames):
    dst.mkdir(parents=True, exist_ok=True)
    for f in FILES.values():
        (dst / f"{f}.csv").write_bytes((SAMPLE / f"{f}.csv").read_bytes())
    for name, df in frames.items():
        df.to_csv(dst / f"{name}.csv", index=False)
    return dst


# ---- 1. detection: what each file is, how it was read, how many rows ------------------------

def _write(fmt, tmp: Path):
    """One file per format; returns (path, part, expected table, expected rows, text expected in read_as)."""
    im, o, r, inv = (_raw(n) for n in ("item_master", "orders", "receipts", "inventory_snapshots"))
    if fmt == "csv":
        return SAMPLE / "orders.csv", None, "orders", len(o), "comma"
    if fmt == "tab":
        p = tmp / "items.txt"; im.to_csv(p, sep="\t", index=False); return p, None, "item_master", len(im), "tab"
    if fmt == "pipe":
        p = tmp / "recv.psv"; r.to_csv(p, sep="|", index=False); return p, None, "receipts", len(r), "pipe"
    if fmt == "utf16":
        p = tmp / "inv.txt"; inv.to_csv(p, sep="\t", index=False, encoding="utf-16")
        return p, None, "inventory", len(inv), "UTF-16"
    if fmt == "gz":
        p = tmp / "orders.csv.gz"
        with gzip.open(p, "wt", encoding="utf-8", newline="") as f:
            o.to_csv(f, index=False)
        return p, None, "orders", len(o), "gzip"
    if fmt == "zip":
        p = tmp / "export.zip"
        with zipfile.ZipFile(p, "w") as z:
            z.write(SAMPLE / "receipts.csv", "data/receipts.csv")
        return p, "data/receipts.csv", "receipts", len(r), "zip"
    if fmt == "xlsx":
        p = tmp / "export.xlsx"
        with pd.ExcelWriter(p, engine="openpyxl") as w:
            im.to_excel(w, sheet_name="Items", index=False, startrow=2)
            w.sheets["Items"]["A1"] = "Item export"
        return p, "Items", "item_master", len(im), "Excel sheet 'Items'"
    if fmt in ("xls", "xlsb", "xlsm"):
        p = FIX / f"excel_sample.{fmt}"
        if not p.exists():
            pytest.skip(f"{p.name} missing")
        n = len(pd.read_excel(FIX / "excel_sample.xlsx", sheet_name="Orders"))
        return p, "Orders", "orders", n, "Excel sheet 'Orders'"
    if fmt == "json":
        p = tmp / "recv.json"; p.write_text(r.to_json(orient="records"), encoding="utf-8")
        return p, None, "receipts", len(r), "JSON"
    if fmt == "jsonl":
        p = tmp / "recv.jsonl"; p.write_text(r.to_json(orient="records", lines=True), encoding="utf-8")
        return p, None, "receipts", len(r), "JSON"
    if fmt == "parquet":
        p = tmp / "items.parquet"
        duckdb.sql(f"copy (select * from read_csv('{(SAMPLE / 'item_master.csv').as_posix()}', header=true)) "
                   f"to '{p.as_posix()}' (format parquet)")
        return p, None, "item_master", len(im), "Parquet"
    if fmt in ("feather", "orc"):
        import pyarrow as pa
        import pyarrow.feather as feather
        import pyarrow.orc as orc
        t = pa.Table.from_pandas(_typed("inventory_snapshots"), preserve_index=False)
        p = tmp / f"inv.{fmt}"
        (orc.write_table if fmt == "orc" else feather.write_feather)(t, p)
        return p, None, "inventory", len(inv), "ORC" if fmt == "orc" else "Arrow"
    if fmt == "sqlite":
        p = tmp / "wms.sqlite"
        with sqlite3.connect(p) as con:
            _typed("orders").to_sql("LINES", con, index=False)
        return p, "LINES", "orders", len(o), "SQLite table 'LINES'"
    if fmt == "duckdb":
        p = tmp / "wms.duckdb"
        con = duckdb.connect(str(p))
        con.execute(f"create table stock as select * from read_csv('{(SAMPLE / 'inventory_snapshots.csv').as_posix()}')")
        con.close()
        return p, "stock", "inventory", len(inv), "DuckDB table 'stock'"
    raise ValueError(fmt)


FORMATS = ["csv", "tab", "pipe", "utf16", "gz", "zip", "xlsx", "xls", "xlsb", "xlsm", "json", "jsonl", "parquet",
           "feather", "orc", "sqlite", "duckdb"]


@pytest.mark.parametrize("fmt", FORMATS)
def test_detection_per_format(tmp_path, fmt):
    path, part, table, rows, read_as = _write(fmt, tmp_path)
    found = {d.part: d for d in detect([path]) if d.file == path}
    d = found[part]
    assert d.table == table and d.confidence in ("high", "medium"), (d.table, d.confidence, d.reason)
    assert d.rows == rows
    assert read_as in d.read_as, d.read_as
    assert not d.at_excel_limit


def test_detection_lists_non_data_files(tmp_path):
    (tmp_path / "notes.docx").write_bytes(b"PK\x03\x04" + b"\x00" * 64)  # a zip that is not data
    (tmp_path / "readme.md").write_text("# hello", encoding="utf-8")
    (tmp_path / "contacts.csv").write_text("Name,Phone\nAna,555\n", encoding="utf-8")
    got = {d.file.name: d for d in detect(sorted(tmp_path.iterdir()))}
    assert set(got) == {"notes.docx", "readme.md", "contacts.csv"}
    assert all(d.table is None for d in got.values())
    assert got["contacts.csv"].rows == 1 and got["notes.docx"].rows is None


def test_detection_flags_excel_row_limit(tmp_path):
    p = tmp_path / "orders.csv"
    _raw("orders").head(500).to_csv(p, index=False)
    d = detect([p], IngestParams(excel_row_limits=(500,)))[0]
    assert d.rows == 500 and d.at_excel_limit


# ---- 2. coverage ------------------------------------------------------------------------------

def _lines(dates):
    dates = pd.to_datetime(pd.Series(list(dates)))
    return pd.DataFrame({"order_id": [f"O{i}" for i in range(len(dates))], "order_date": dates,
                         "sku": "A", "qty": 1})


def test_coverage_full_year_with_peak():
    c = coverage(_lines(pd.date_range("2025-01-01", "2025-12-31")), IngestParams())
    assert c.span_days == 365 and c.days_with_rows == 365 and not c.gap_runs
    assert c.peak_covered and not c.short and not c.warnings


def test_coverage_twelve_months_missing_peak():
    days = pd.date_range("2024-10-01", "2025-09-30")
    days = days[~days.month.isin([11, 12])]
    c = coverage(_lines(days), IngestParams())
    assert not c.short and not c.peak_covered
    assert any("Nov" in w for w in c.warnings)
    assert c.gap_runs[0][0] == pd.Timestamp("2024-11-01") and c.gap_runs[0][2] == 61


def test_coverage_lists_known_gaps():
    days = pd.date_range("2025-01-01", "2025-12-31")
    days = days[~((days >= "2025-03-03") & (days <= "2025-03-09"))]
    c = coverage(_lines(days), IngestParams())
    assert (pd.Timestamp("2025-03-03"), pd.Timestamp("2025-03-09"), 7) in c.gap_runs
    assert c.daily.loc[c.daily["date"] == "2025-03-05", "lines"].iat[0] == 0


def test_sample_client_is_short_and_misses_peak():
    r = ingest_folder(SAMPLE, profile=load_wms_profile(WMS))
    codes = set(r.issues.loc[r.issues.severity == "warn", "code"])
    assert {"coverage_short", "peak_not_covered"} <= codes
    assert r.coverage.first == pd.Timestamp("2024-12-15") and r.coverage.short and not r.coverage.peak_covered
    assert r.coverage.lines == len(r.tables["orders"])


def test_clipped_export_blocks_and_coverage_ends_early(tmp_path):
    o = _raw("orders")
    o["_d"] = pd.to_datetime(o["Order Date"], format="%m/%d/%Y")
    clipped = o.sort_values("_d", kind="stable").head(6000).drop(columns="_d")
    folder = _clone(tmp_path / "c", orders=clipped)
    r = ingest_folder(folder, profile=load_wms_profile(WMS), params=IngestParams(excel_row_limits=(6000,)),
                      strict=False)
    assert "excel_row_limit" in set(r.issues.loc[r.issues.severity == "block", "code"])
    assert r.coverage.last < pd.Timestamp("2025-04-13")


# ---- 3. plain language for every issue code ----------------------------------------------------

def _emitted_codes():
    src = (ROOT / "src" / "gtpfit" / "ingest" / "pipeline.py").read_text(encoding="utf-8")
    codes = set(re.findall(r'add\((?:"", |table, )?"([a-z_]+)", "(?:block|warn|info)"', src))
    codes |= set(re.findall(r'Issue\("[a-z_]+", "([a-z_]+)", "(?:block|warn|info)"', src))
    codes |= {"bad_qty", "missing_value", "unparseable", "coverage_short", "peak_not_covered"}  # built in loops
    return codes


def test_every_issue_code_has_plain_language():
    codes = _emitted_codes()
    assert len(codes) > 35
    missing = sorted(codes - set(ISSUE_TEXT))
    assert not missing, missing
    for code in codes:
        t = ISSUE_TEXT[code]
        assert t.title and t.what and t.impact, code


# ---- 4. accepting a block needs a reason -------------------------------------------------------

def _gap_folder(tmp_path):
    o = _raw("orders")
    d = pd.to_datetime(o["Order Date"], format="%m/%d/%Y")
    return _clone(tmp_path / "g", orders=o[~d.between("2025-02-03", "2025-02-12")])


def test_accept_requires_a_reason(tmp_path):
    folder = _gap_folder(tmp_path)
    for bad in ({"date_gap": "  "}, ["date_gap"]):
        with pytest.raises(ValueError, match="reason"):
            ingest_folder(folder, profile=load_wms_profile(WMS), accept=bad, strict=False)


def test_accept_is_per_table_and_recorded(tmp_path):
    folder = _gap_folder(tmp_path)
    wrong = ingest_folder(folder, profile=load_wms_profile(WMS), accept={"receipts:date_gap": "n/a"}, strict=False)
    assert wrong.blocked
    r = ingest_folder(folder, profile=load_wms_profile(WMS),
                      accept={"orders:date_gap": {"reason": "Site closed for stocktake", "at": "2026-10-21T09:00"}})
    assert not r.blocked
    acc = r.accepted
    assert len(acc) == 1 and acc[0]["table"] == "orders" and acc[0]["code"] == "date_gap"
    assert acc[0]["reason"] == "Site closed for stocktake" and acc[0]["accepted_at"] == "2026-10-21T09:00"
    out = write_outputs(r, tmp_path / "_ingest", client="Acme")
    ab = pd.read_csv(out / "accepted_blocks.csv")
    assert list(ab["reason"]) == ["Site closed for stocktake"] and list(ab["code"]) == ["date_gap"]


def test_cli_accept_needs_code_equals_reason(tmp_path, capsys):
    from gtpfit.ingest.cli import main
    folder = _gap_folder(tmp_path)
    with pytest.raises(SystemExit) as e:
        main([str(folder), "--wms", str(WMS), "--accept", "date_gap", "--out", str(tmp_path / "o1")])
    assert e.value.code == 2 and "reason" in capsys.readouterr().out.lower()
    main([str(folder), "--wms", str(WMS), "--accept", "orders:date_gap=closed for stocktake",
          "--out", str(tmp_path / "o2")])
    assert pd.read_csv(tmp_path / "o2" / "accepted_blocks.csv")["reason"].iat[0] == "closed for stocktake"


# ---- 5. progress --------------------------------------------------------------------------------

def test_progress_reports_each_table():
    from gtpfit.ingest import discover, ingest_sources
    events = []
    ingest_sources(discover(sorted(SAMPLE.iterdir())), profile=load_wms_profile(WMS), strict=False,
                   on_progress=events.append)
    starts = [e["table"] for e in events if e["stage"] == "start"]
    assert starts == ["item_master", "orders", "receipts", "inventory"]
    done = {e["table"]: e["rows"] for e in events if e["stage"] == "done"}
    assert done["orders"] > 20000 and all(e["total"] == 4 for e in events)


# ---- 6. client-facing outputs -------------------------------------------------------------------

def test_outputs_receipt_and_rejects(tmp_path):
    r = ingest_folder(SAMPLE, profile=load_wms_profile(WMS))
    out = write_outputs(r, tmp_path / "_ingest", client="Acme Retail")
    html = (out / "data_receipt.html").read_text(encoding="utf-8")
    assert "Acme Retail" in html and "orders.csv" in html and f"{len(r.tables['orders']):,}" in html
    assert not re.search(r'(src|href)="https?://', html)  # self-contained
    rej = pd.read_csv(out / "rejects_for_client.csv", dtype=str)
    assert {"record", "table", "file", "row", "reason"} <= set(rej.columns)
    rows = rej[rej["record"] == "rejected row"]
    assert len(rows) == len(r.rejects["orders"]) and set(rows["reason"]) == {"bad_qty"}
    assert rows["file"].eq("orders.csv").all() and rows["row"].astype(int).min() >= 2
    assert (out / "accepted_blocks.csv").exists()


def test_data_request_template_lists_required_fields(tmp_path):
    from openpyxl import load_workbook
    from gtpfit.ingest.request_template import request_html, request_xlsx
    from gtpfit.ingest.schema import FIELD_HELP, TABLES
    html = request_html()
    assert "12 months" in html and "November" in html
    for table, flds in TABLES.items():
        for f in flds:
            if f.required:
                assert FIELD_HELP[table][f.name].split(":")[0] in html, (table, f.name)
    p = tmp_path / "request.xlsx"
    p.write_bytes(request_xlsx())
    ws = load_workbook(p).active
    cells = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Unit weight" in cells and "Required" in cells


# ---- 7. a folder read in place: one-line summary, folder in the receipt ------------------------------

def test_folder_summary_counts(tmp_path):
    from gtpfit.ingest.detect import summarize
    folder = _clone(tmp_path / "client")
    (folder / "notes.md").write_text("# notes", encoding="utf-8")
    (folder / "contacts.csv").write_text("Name,Phone\nAna,555\n", encoding="utf-8")
    with zipfile.ZipFile(folder / "extra.zip", "w") as z:
        z.write(SAMPLE / "receipts.csv", "receipts.csv")
    files = sorted(p for p in folder.iterdir() if p.is_file())
    s = summarize(files, detect(files))
    assert (s.files, s.data, s.ignored, s.zips) == (7, 4, 2, 1)
    assert s.text == "7 files found: 4 data, 2 ignored, 1 zip"


def test_folder_summary_text_without_zip_or_ignored(tmp_path):
    from gtpfit.ingest.detect import summarize
    files = sorted(p for p in _clone(tmp_path / "c").iterdir() if p.is_file())
    assert summarize(files, detect(files)).text == "4 files found: 4 data, 0 ignored, 0 zips"


def test_receipt_records_the_folder(tmp_path):
    r = ingest_folder(SAMPLE, profile=load_wms_profile(WMS))
    out = write_outputs(r, tmp_path / "_ingest", client="Acme", folder=SAMPLE)
    html = (out / "data_receipt.html").read_text(encoding="utf-8")
    assert "Folder read" in html and str(SAMPLE) in html


# ---- busiest-months hint (a hint only; never sets peak_months) --------------------------------------------

def _dates(start, end, per_day):
    """One date per order line: per_day(day) lines on each calendar day."""
    return pd.Series([d for d in pd.date_range(start, end) for _ in range(per_day(d))])


def test_peak_hint_single_peak():
    from gtpfit.ingest.coverage import busiest_months
    h = busiest_months(_dates("2025-01-01", "2025-12-31", lambda d: 4 if d.month == 12 else 2), (11, 12))
    assert h.enough and h.months == (12,) and h.ratios[12] == pytest.approx(2.0)
    assert h.text == "Busiest months in this data: Dec (2.0×)." and not h.outside


def test_peak_hint_two_peaks_outside_the_peak_months():
    from gtpfit.ingest.coverage import OUTSIDE_NOTE, busiest_months
    h = busiest_months(_dates("2025-01-01", "2025-12-31", lambda d: 3 if d.month in (5, 6) else 2), (11, 12))
    assert h.months == (5, 6) and h.outside
    assert "Worth asking the client" in OUTSIDE_NOTE
    assert not busiest_months(_dates("2025-01-01", "2025-12-31", lambda d: 3 if d.month in (5, 6) else 2),
                              (6, 12)).outside  # one hinted month is set: no note


def test_peak_hint_flat_and_short_history():
    from gtpfit.ingest.coverage import busiest_months
    flat = busiest_months(_dates("2025-01-01", "2025-12-31", lambda d: 2), (11, 12))
    assert flat.enough and flat.months == () and not flat.outside
    assert flat.text == "Busiest months in this data: no month stands out."
    short = busiest_months(_dates("2025-01-01", "2025-11-30", lambda d: 5 if d.month == 11 else 1), (5,))
    assert not short.enough and short.months == () and not short.outside
    assert "not enough history to suggest" in short.text


def test_peak_hint_partial_months_and_years_averaged():
    from gtpfit.ingest.coverage import busiest_months
    # history starting mid-month: a partial month is not made to look quiet (lines per covered day)
    h = busiest_months(_dates("2024-12-20", "2025-12-19", lambda d: 2), ())
    assert h.months == () and set(h.ratios.values()) == {1.0}
    # two years; December at 1.5x in one year only averages to 1.25x: under the 1.3 threshold
    two = busiest_months(_dates("2024-01-01", "2025-12-31",
                                lambda d: 3 if (d.month == 12 and d.year == 2024) else 2), (11, 12))
    assert two.ratios[12] == pytest.approx(1.25) and two.months == ()
    # a December spike in both years is named
    both = busiest_months(_dates("2024-01-01", "2025-12-31", lambda d: 3 if d.month == 12 else 2), (11, 12))
    assert both.months == (12,)


def test_peak_hint_threshold_is_a_param():
    from gtpfit.ingest.coverage import busiest_months
    data = _dates("2025-01-01", "2025-12-31", lambda d: 5 if d.month == 7 else 4)  # July at 1.25x
    assert busiest_months(data, (), IngestParams().peak_hint_ratio).months == ()
    assert busiest_months(data, (), 1.2).months == (7,)


# ---- re-running cheaply: acceptances after the run, per-table inputs ----------------------------------------

def test_apply_acceptances_matches_accepting_during_the_run(tmp_path):
    from gtpfit.ingest import apply_acceptances
    o = _raw("orders")
    o["_d"] = pd.to_datetime(o["Order Date"], format="%m/%d/%Y")
    folder = _clone(tmp_path / "c", orders=o.sort_values("_d", kind="stable").head(6000).drop(columns="_d"))
    params = IngestParams(excel_row_limits=(6000,))
    raw = ingest_folder(folder, profile=load_wms_profile(WMS), params=params, strict=False)
    key = "orders:excel_row_limit"
    assert raw.blocked
    accept = {key: "client confirmed the extract"}
    after = apply_acceptances(raw, accept)
    during = ingest_folder(folder, profile=load_wms_profile(WMS), params=params, accept=accept, strict=False)
    cols = ["table", "code", "severity", "message"]
    assert after.issues[cols].equals(during.issues[cols])
    assert not after.blocked and [a["reason"] for a in after.accepted] == ["client confirmed the extract"]
    assert raw.blocked  # the cached run itself is left as it was
    undone = apply_acceptances(after, {})
    assert undone.blocked and undone.issues[cols].equals(raw.issues[cols])


def test_table_inputs_units_only_reach_the_item_master():
    from gtpfit.ingest import table_inputs
    from gtpfit.ingest.mapping import WmsProfile
    a = WmsProfile(name="x", units={"length": "in", "weight": "lb"})
    b = WmsProfile(name="x", units={"length": "cm", "weight": "kg"})
    assert table_inputs("item_master", a) != table_inputs("item_master", b)
    for t in ("orders", "receipts", "inventory"):
        assert table_inputs(t, a) == table_inputs(t, b)
    # thresholds only the coverage step reads leave every table's inputs alone; one ingest_table reads does not
    p = IngestParams()
    assert table_inputs("orders", a, p) == table_inputs("orders", a, IngestParams(peak_months=(6,)))
    assert table_inputs("orders", a, p) != table_inputs("orders", a, IngestParams(max_gap_days_orders=9))
